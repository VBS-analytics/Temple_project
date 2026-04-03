from datetime import date, datetime, timedelta
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch

from django.db.models import F, Window
from django.db.models.functions import RowNumber
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from rest_framework import status
from rest_framework.test import APIClient

from accounts.access import (
    EXPENSE_TRACKER_ACCESS_DENIED_MESSAGE,
    REPORT_DOWNLOAD_ACCESS_DENIED_MESSAGE,
)
from accounts.models import DonorProfile, User, UserRole
from pooja.models import (
    DayOptionCategory,
    PoojaDayOption,
    PoojaOption,
    PoojaRegistration,
    RecurrenceFrequency,
    RecurrenceKind,
    RecurringPoojaPlan,
)
from pooja.services.recurrence import _clean_stale_chrt_dues
from payments.models import PassbookEntry
from payments.models import (
    AdditionIncomeRecord,
    CombinePaymentMapping,
    Donation,
    ExpenseCategory,
    ExpenseRecord,
    IncomeCategory,
    PaymentRecord,
    PaymentStatus,
)
from payments.services import regenerate_donor_passbook
from payments.views import _donor_passbook_needs_refresh


class DonationApiTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.create_url = reverse("donation-create")

    def test_anyone_can_create_donation_without_login(self):
        response = self.client.post(
            self.create_url,
            {
                "donor_name": "Lakshmi V",
                "donor_phone_no": "+919876543210",
                "transaction_id": "DON-1001",
                "amount_paid": "750.00",
                "donation_date": "2026-02-03",
                "notes": "Monthly donation",
            },
            format="json",
        )
        self.assertEqual(response.status_code, 201)
        created = Donation.objects.get(id=response.json()["id"])
        self.assertEqual(created.donor_name, "Lakshmi V")
        self.assertEqual(created.donor_phone_no, "+919876543210")
        self.assertEqual(created.transaction_id, "DON-1001")
        self.assertEqual(created.amount_paid, Decimal("750.00"))

    def test_missing_required_fields_returns_400(self):
        response = self.client.post(
            self.create_url,
            {
                "donor_name": "Lakshmi V",
                "amount_paid": "750.00",
                "donation_date": "2026-02-03",
            },
            format="json",
        )
        self.assertEqual(response.status_code, 400)


class PaymentExportAccessTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.full_admin = User.objects.create_superuser(
            phone_number="9999999999",
            name="Full Admin",
            password="adminpass",
        )
        self.restricted_admin = User.objects.create_superuser(
            phone_number="+91 9999999997",
            name="Restricted Admin",
            password="adminpass2",
        )

    def test_restricted_admin_cannot_download_payment_details_export(self):
        self.client.force_authenticate(self.restricted_admin)
        response = self.client.get(reverse("payment-details-export"))
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)
        self.assertEqual(response.json()["detail"], REPORT_DOWNLOAD_ACCESS_DENIED_MESSAGE)

    def test_restricted_admin_cannot_download_general_donation_export(self):
        self.client.force_authenticate(self.restricted_admin)
        response = self.client.get(reverse("general-donation-export"))
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)
        self.assertEqual(response.json()["detail"], REPORT_DOWNLOAD_ACCESS_DENIED_MESSAGE)

    def test_full_admin_can_download_payment_details_export(self):
        self.client.force_authenticate(self.full_admin)
        response = self.client.get(reverse("payment-details-export"))
        self.assertEqual(response.status_code, status.HTTP_200_OK)


class PaymentDetailsExportContentTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.admin = User.objects.create_superuser(
            phone_number="9444444444",
            name="Export Admin",
            password="adminpass",
        )
        self.main_donor = User.objects.create_user(
            phone_number="+919111111111",
            name="Main Donor",
            password="secret",
        )
        self.subordinate_donor = User.objects.create_user(
            phone_number="+919222222222",
            name="Subordinate Donor",
            password="secret",
        )
        self.regular_donor = User.objects.create_user(
            phone_number="+919333333333",
            name="Regular Donor",
            password="secret",
        )

        month_start = timezone.localdate().replace(day=1)
        CombinePaymentMapping.objects.create(
            main_donor=self.main_donor,
            parent_donor=self.subordinate_donor,
            effective_from=month_start,
        )

        self.main_payment = self._create_payment(self.main_donor, "MAIN-TXN-1", Decimal("600.00"))
        self.subordinate_payment = self._create_payment(
            self.subordinate_donor, "SUB-TXN-1", Decimal("500.00")
        )
        self.regular_payment = self._create_payment(self.regular_donor, "REG-TXN-1", Decimal("400.00"))

        self._create_passbook_entry(self.main_donor, self.main_payment, Decimal("600.00"))
        self._create_passbook_entry(self.subordinate_donor, self.subordinate_payment, Decimal("500.00"))
        self._create_passbook_entry(self.regular_donor, self.regular_payment, Decimal("400.00"))

    def _create_payment(self, donor: User, transaction_reference: str, amount: Decimal) -> PaymentRecord:
        return PaymentRecord.objects.create(
            donor=donor,
            amount=amount,
            currency="INR",
            mode="upi",
            status=PaymentStatus.SUCCESS,
            transaction_reference=transaction_reference,
            payment_month=timezone.localdate().replace(day=1),
        )

    def _create_passbook_entry(
        self,
        donor: User,
        payment_record: PaymentRecord,
        paid_amount: Decimal,
    ) -> PassbookEntry:
        return PassbookEntry.objects.create(
            donor=donor,
            entry_date=timezone.localdate().replace(day=1),
            entry_type="paid",
            transaction_details=payment_record.transaction_reference,
            payment_record=payment_record,
            opening_balance=Decimal("0.00"),
            due_amount=paid_amount,
            paid_amount=paid_amount,
            closing_due=Decimal("0.00"),
        )

    def test_payment_details_export_excludes_active_subordinate_donors_from_all_sheets(self):
        self.client.force_authenticate(self.admin)

        with patch("payments.views.regenerate_all_passbooks"):
            response = self.client.get(reverse("payment-details-export"))

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        payload = b"".join(response.streaming_content)
        workbook = load_workbook(filename=BytesIO(payload))

        payment_rows = list(workbook["Payment Records"].iter_rows(min_row=2, values_only=True))
        passbook_rows = list(workbook["Passbook Entries"].iter_rows(min_row=2, values_only=True))
        statement_rows = list(workbook["Donor Statements"].iter_rows(min_row=2, values_only=True))

        payment_donor_ids = {row[1] for row in payment_rows if row and row[1] is not None}
        passbook_donor_ids = {row[1] for row in passbook_rows if row and row[1] is not None}
        statement_donor_ids = {row[0] for row in statement_rows if row and row[0] is not None}

        expected_ids = {self.main_donor.id, self.regular_donor.id}
        self.assertSetEqual(payment_donor_ids, expected_ids)
        self.assertSetEqual(passbook_donor_ids, expected_ids)
        self.assertSetEqual(statement_donor_ids, expected_ids)
        self.assertNotIn(self.subordinate_donor.id, payment_donor_ids)
        self.assertNotIn(self.subordinate_donor.id, passbook_donor_ids)
        self.assertNotIn(self.subordinate_donor.id, statement_donor_ids)


class PassbookRefreshDetectionTests(TestCase):
    def setUp(self):
        self.donor = User.objects.create_user(
            phone_number="+919000000001",
            name="Refresh Donor",
            password="secret",
        )
        self.pooja_option = PoojaOption.objects.create(code="R1", name="Recurring 1")
        self.day_option = PoojaDayOption.objects.create(
            code="REGULAR_TEST",
            description="Regular",
            category=DayOptionCategory.CODE,
        )

    def test_needs_refresh_when_plan_changes_after_passbook_generation(self):
        RecurringPoojaPlan.objects.create(
            donor=self.donor,
            pooja_option=self.pooja_option,
            day_option=self.day_option,
            recurrence_kind=RecurrenceKind.RECURRING,
            recurrence_frequency=RecurrenceFrequency.MONTHLY,
            start_date=date(2026, 1, 1),
            next_occurrence=date(2026, 1, 1),
            amount=Decimal("500.00"),
            is_active=True,
        )

        with patch("payments.services.timezone.localdate", return_value=date(2026, 2, 11)):
            regenerate_donor_passbook(self.donor.id)

        current_month = date(2026, 2, 1)
        self.assertFalse(_donor_passbook_needs_refresh(self.donor.id, current_month))

        RecurringPoojaPlan.objects.create(
            donor=self.donor,
            pooja_option=PoojaOption.objects.create(code="R2", name="Recurring 2"),
            day_option=self.day_option,
            recurrence_kind=RecurrenceKind.RECURRING,
            recurrence_frequency=RecurrenceFrequency.MONTHLY,
            start_date=date(2026, 1, 1),
            next_occurrence=date(2026, 1, 1),
            amount=Decimal("100.00"),
            is_active=True,
        )

        self.assertTrue(_donor_passbook_needs_refresh(self.donor.id, current_month))

    def test_needs_refresh_when_earliest_due_month_is_missing(self):
        RecurringPoojaPlan.objects.create(
            donor=self.donor,
            pooja_option=self.pooja_option,
            day_option=self.day_option,
            recurrence_kind=RecurrenceKind.RECURRING,
            recurrence_frequency=RecurrenceFrequency.MONTHLY,
            start_date=date(2026, 1, 1),
            next_occurrence=date(2026, 1, 1),
            amount=Decimal("500.00"),
            is_active=True,
        )

        # Simulate stale passbook generated with old logic: Jan due missing, Feb due present.
        PassbookEntry.objects.create(
            donor=self.donor,
            entry_date=date(2025, 12, 31),
            entry_type="balance",
            transaction_details="-",
            opening_balance=Decimal("0.00"),
            due_amount=Decimal("0.00"),
            paid_amount=Decimal("0.00"),
            closing_due=Decimal("0.00"),
        )
        PassbookEntry.objects.create(
            donor=self.donor,
            entry_date=date(2026, 2, 1),
            entry_type="due",
            transaction_details="--- Pooja DUE ---",
            opening_balance=Decimal("0.00"),
            due_amount=Decimal("500.00"),
            paid_amount=Decimal("0.00"),
            closing_due=Decimal("500.00"),
        )

        current_month = date(2026, 2, 1)
        self.assertTrue(_donor_passbook_needs_refresh(self.donor.id, current_month))

    def test_needs_refresh_when_monthly_due_exists_before_anchor_month(self):
        RecurringPoojaPlan.objects.create(
            donor=self.donor,
            pooja_option=self.pooja_option,
            day_option=self.day_option,
            recurrence_kind=RecurrenceKind.RECURRING,
            recurrence_frequency=RecurrenceFrequency.MONTHLY,
            start_date=date(2026, 3, 1),
            next_occurrence=date(2026, 3, 1),
            amount=Decimal("500.00"),
            is_active=True,
        )

        # Simulate stale passbook generated with old logic: Feb due exists
        # even though recurring anchor starts in Mar.
        PassbookEntry.objects.create(
            donor=self.donor,
            entry_date=date(2025, 12, 31),
            entry_type="balance",
            transaction_details="-",
            opening_balance=Decimal("0.00"),
            due_amount=Decimal("0.00"),
            paid_amount=Decimal("0.00"),
            closing_due=Decimal("0.00"),
        )
        PassbookEntry.objects.create(
            donor=self.donor,
            entry_date=date(2026, 2, 1),
            entry_type="due",
            transaction_details="--- Pooja DUE ---",
            opening_balance=Decimal("0.00"),
            due_amount=Decimal("500.00"),
            paid_amount=Decimal("0.00"),
            closing_due=Decimal("500.00"),
        )

        current_month = date(2026, 2, 1)
        self.assertTrue(_donor_passbook_needs_refresh(self.donor.id, current_month))


class PassbookDeduplicationTests(TestCase):
    def test_dedup_prefers_latest_created_entry(self):
        donor = User.objects.create_user(
            phone_number="+919000000002",
            name="Dedup Donor",
            password="secret",
        )

        older = PassbookEntry.objects.create(
            donor=donor,
            entry_date=date(2026, 2, 1),
            entry_type="due",
            transaction_details="--- Pooja DUE ---",
            opening_balance=Decimal("1800.00"),
            due_amount=Decimal("500.00"),
            paid_amount=Decimal("0.00"),
            closing_due=Decimal("2300.00"),
        )
        newer = PassbookEntry.objects.create(
            donor=donor,
            entry_date=date(2026, 2, 1),
            entry_type="due",
            transaction_details="--- Pooja DUE ---",
            opening_balance=Decimal("1800.00"),
            due_amount=Decimal("600.00"),
            paid_amount=Decimal("0.00"),
            closing_due=Decimal("2400.00"),
        )

        now = timezone.now()
        PassbookEntry.objects.filter(pk=older.pk).update(created_at=now - timedelta(minutes=5))
        PassbookEntry.objects.filter(pk=newer.pk).update(created_at=now)

        row = (
            PassbookEntry.objects.filter(donor=donor)
            .annotate(
                rn=Window(
                    expression=RowNumber(),
                    partition_by=[F("donor_id"), F("entry_date"), F("entry_type")],
                    order_by=F("created_at").desc(),
                )
            )
            .filter(rn=1)
            .get()
        )
        self.assertEqual(row.due_amount, Decimal("600.00"))

    def test_dedup_keeps_distinct_paid_rows_on_same_date(self):
        donor = User.objects.create_user(
            phone_number="+919000000003",
            name="Paid Dedup Donor",
            password="secret",
        )

        first_payment = PaymentRecord.objects.create(
            donor=donor,
            amount=Decimal("500.00"),
            mode="upi",
            status=PaymentStatus.SUCCESS,
            transaction_reference="TXN-A",
            payment_month=date(2026, 2, 1),
        )
        second_payment = PaymentRecord.objects.create(
            donor=donor,
            amount=Decimal("600.00"),
            mode="upi",
            status=PaymentStatus.SUCCESS,
            transaction_reference="TXN-B",
            payment_month=date(2026, 2, 1),
        )

        PassbookEntry.objects.create(
            donor=donor,
            entry_date=date(2026, 2, 1),
            entry_type="paid",
            transaction_details="TXN-A",
            payment_record=first_payment,
            opening_balance=Decimal("1000.00"),
            due_amount=Decimal("0.00"),
            paid_amount=Decimal("500.00"),
            closing_due=Decimal("500.00"),
        )
        PassbookEntry.objects.create(
            donor=donor,
            entry_date=date(2026, 2, 1),
            entry_type="paid",
            transaction_details="TXN-B",
            payment_record=second_payment,
            opening_balance=Decimal("500.00"),
            due_amount=Decimal("0.00"),
            paid_amount=Decimal("600.00"),
            closing_due=Decimal("-100.00"),
        )

        rows = (
            PassbookEntry.objects.filter(donor=donor)
            .annotate(
                rn=Window(
                    expression=RowNumber(),
                    partition_by=[
                        F("donor_id"),
                        F("entry_date"),
                        F("entry_type"),
                        F("payment_record_id"),
                        F("registration_id"),
                        F("transaction_details"),
                        F("due_amount"),
                        F("paid_amount"),
                    ],
                    order_by=F("created_at").desc(),
                )
            )
            .filter(rn=1)
        )

        self.assertEqual(rows.count(), 2)


class PassbookOrderingTests(TestCase):
    def test_same_day_due_and_paid_apply_due_before_paid(self):
        donor = User.objects.create_user(
            phone_number="+919000000004",
            name="Ordering Donor",
            password="secret",
        )
        pooja_option = PoojaOption.objects.create(code="R3", name="Recurring 3")
        day_option = PoojaDayOption.objects.create(
            code="REGULAR_ORDER_TEST",
            description="Regular",
            category=DayOptionCategory.CODE,
        )
        RecurringPoojaPlan.objects.create(
            donor=donor,
            pooja_option=pooja_option,
            day_option=day_option,
            recurrence_kind=RecurrenceKind.RECURRING,
            recurrence_frequency=RecurrenceFrequency.MONTHLY,
            start_date=date(2026, 1, 1),
            next_occurrence=date(2026, 1, 1),
            amount=Decimal("1000.00"),
            is_active=True,
        )
        PaymentRecord.objects.create(
            donor=donor,
            amount=Decimal("1000.00"),
            mode="upi",
            status=PaymentStatus.SUCCESS,
            transaction_reference="FA19147194",
            payment_month=date(2026, 1, 1),
        )

        with patch("payments.services.timezone.localdate", return_value=date(2026, 2, 11)):
            regenerate_donor_passbook(donor.id)

        jan_entries = list(
            PassbookEntry.objects.filter(donor=donor, entry_date=date(2026, 1, 1))
            .exclude(entry_type="balance")
            .order_by("id")
        )
        self.assertEqual(len(jan_entries), 2)
        self.assertEqual(jan_entries[0].entry_type, "due")
        self.assertEqual(jan_entries[0].closing_due, Decimal("1000.00"))
        self.assertEqual(jan_entries[1].entry_type, "paid")
        self.assertEqual(jan_entries[1].closing_due, Decimal("0.00"))


class PassbookCancellationFreezeTests(TestCase):
    def test_cancel_from_month_freezes_older_month_due_backfill(self):
        donor = User.objects.create_user(
            phone_number="+919000000005",
            name="Cancel Freeze Donor",
            password="secret",
        )
        pooja_option = PoojaOption.objects.create(code="R4", name="Recurring 4")
        day_option = PoojaDayOption.objects.create(
            code="REGULAR_CANCEL_TEST",
            description="Regular",
            category=DayOptionCategory.CODE,
        )

        # Active baseline plan
        RecurringPoojaPlan.objects.create(
            donor=donor,
            pooja_option=pooja_option,
            day_option=day_option,
            recurrence_kind=RecurrenceKind.RECURRING,
            recurrence_frequency=RecurrenceFrequency.MONTHLY,
            start_date=date(2026, 1, 1),
            next_occurrence=date(2026, 1, 1),
            amount=Decimal("500.00"),
            is_active=True,
        )

        # Canceled plan effective from Apr 2026 (historical months before Apr must stay as-is)
        RecurringPoojaPlan.objects.create(
            donor=donor,
            pooja_option=PoojaOption.objects.create(code="R5", name="Recurring 5"),
            day_option=day_option,
            recurrence_kind=RecurrenceKind.RECURRING,
            recurrence_frequency=RecurrenceFrequency.MONTHLY,
            start_date=date(2026, 1, 1),
            next_occurrence=None,
            amount=Decimal("2500.00"),
            is_active=False,
            pause_from=date(2026, 4, 1),
            pause_until=date.max,
            metadata={
                "canceled_at": "2026-04-02",
                "cancel_effective_from": "2026-04-01",
            },
        )

        # Existing historical dues already stored in records.
        # Jan must remain ₹500, even though active-plan total for Jan would be ₹3000.
        for payment_month, amount in (
            (date(2026, 1, 1), Decimal("500.00")),
            (date(2026, 2, 1), Decimal("3000.00")),
            (date(2026, 3, 1), Decimal("3000.00")),
            (date(2026, 4, 1), Decimal("500.00")),
        ):
            PaymentRecord.objects.create(
                donor=donor,
                amount=amount,
                mode="pending",
                status=PaymentStatus.PENDING,
                payment_month=payment_month,
                notes="Monthly recurring pooja contribution due",
                registration=None,
            )

        with patch("payments.services.timezone.localdate", return_value=date(2026, 4, 2)):
            regenerate_donor_passbook(donor.id, ensure_dues=False)

        jan_due_entry = (
            PassbookEntry.objects.filter(
                donor=donor,
                entry_date=date(2026, 1, 1),
                entry_type="due",
            )
            .order_by("id")
            .first()
        )

        self.assertIsNotNone(jan_due_entry)
        self.assertEqual(jan_due_entry.due_amount, Decimal("500.00"))

    def test_freeze_still_backfills_missing_historical_month_without_split_due_rows(self):
        donor = User.objects.create_user(
            phone_number="+919000000011",
            name="Freeze Missing Month Donor",
            password="secret",
        )
        day_option = PoojaDayOption.objects.create(
            code="REG_MISSING_MONTH",
            description="Regular",
            category=DayOptionCategory.CODE,
        )

        jan1 = date(2026, 1, 1)
        jan14 = date(2026, 1, 14)
        apr1 = date(2026, 4, 1)
        jan1_dt = timezone.make_aware(datetime(2026, 1, 1, 9, 0, 0))
        jan14_dt = timezone.make_aware(datetime(2026, 1, 14, 9, 0, 0))

        reg_a = PoojaRegistration.objects.create(
            donor=donor,
            pooja_option=PoojaOption.objects.create(code="RM1", name="Recurring Missing 1"),
            day_option=day_option,
            start_date=jan1,
            total_amount=Decimal("500.00"),
        )
        reg_b = PoojaRegistration.objects.create(
            donor=donor,
            pooja_option=PoojaOption.objects.create(code="RM2", name="Recurring Missing 2"),
            day_option=day_option,
            start_date=jan14,
            total_amount=Decimal("1000.00"),
        )
        PoojaRegistration.objects.filter(pk=reg_a.pk).update(created_at=jan1_dt)
        PoojaRegistration.objects.filter(pk=reg_b.pk).update(created_at=jan14_dt)

        RecurringPoojaPlan.objects.create(
            donor=donor,
            pooja_option=reg_a.pooja_option,
            day_option=day_option,
            recurrence_kind=RecurrenceKind.RECURRING,
            recurrence_frequency=RecurrenceFrequency.MONTHLY,
            start_date=jan1,
            next_occurrence=None,
            amount=Decimal("500.00"),
            is_active=False,
            pause_from=apr1,
            pause_until=date.max,
            origin_registration=reg_a,
            metadata={
                "canceled_at": "2026-04-02",
                "cancel_effective_from": "2026-04-01",
            },
        )
        RecurringPoojaPlan.objects.create(
            donor=donor,
            pooja_option=reg_b.pooja_option,
            day_option=day_option,
            recurrence_kind=RecurrenceKind.RECURRING,
            recurrence_frequency=RecurrenceFrequency.MONTHLY,
            start_date=jan14,
            next_occurrence=apr1,
            amount=Decimal("1000.00"),
            is_active=True,
            origin_registration=reg_b,
        )

        # Intentionally keep January without a registration-null pending due row.
        # Legacy donors can still have successful payments plus registration rows.
        PaymentRecord.objects.create(
            donor=donor,
            amount=Decimal("1500.00"),
            mode="upi",
            status=PaymentStatus.SUCCESS,
            transaction_reference="RM-JAN-SUCCESS",
            payment_month=jan1,
        )

        with patch("payments.services.timezone.localdate", return_value=date(2026, 4, 2)):
            regenerate_donor_passbook(donor.id, ensure_dues=False)

        jan_due_entries = list(
            PassbookEntry.objects.filter(
                donor=donor,
                entry_type="due",
                entry_date__year=2026,
                entry_date__month=1,
            ).order_by("entry_date", "id")
        )

        self.assertEqual(len(jan_due_entries), 1)
        self.assertEqual(jan_due_entries[0].entry_date, jan1)
        self.assertEqual(jan_due_entries[0].due_amount, Decimal("1500.00"))


class RecurringDueGenerationEdgeCaseTests(TestCase):
    def test_future_success_payment_month_does_not_block_current_month_due(self):
        donor = User.objects.create_user(
            phone_number="+919000000006",
            name="Future Payment Donor",
            password="secret",
        )
        pooja_option = PoojaOption.objects.create(code="R6", name="Recurring 6")
        day_option = PoojaDayOption.objects.create(
            code="REGFUTPAY",
            description="Regular",
            category=DayOptionCategory.CODE,
        )

        RecurringPoojaPlan.objects.create(
            donor=donor,
            pooja_option=pooja_option,
            day_option=day_option,
            recurrence_kind=RecurrenceKind.RECURRING,
            recurrence_frequency=RecurrenceFrequency.MONTHLY,
            start_date=date(2026, 4, 1),
            next_occurrence=date(2026, 4, 1),
            amount=Decimal("1000.00"),
            is_active=True,
        )

        # Future-month success payment should not suppress April due generation.
        PaymentRecord.objects.create(
            donor=donor,
            amount=Decimal("1000.00"),
            mode="upi",
            status=PaymentStatus.SUCCESS,
            transaction_reference="FUTURE-TXN-1",
            payment_month=date(2026, 5, 1),
        )

        with patch("payments.services.timezone.localdate", return_value=date(2026, 4, 2)):
            regenerate_donor_passbook(donor.id)

        april_due = PaymentRecord.objects.filter(
            donor=donor,
            registration__isnull=True,
            status=PaymentStatus.PENDING,
            payment_month=date(2026, 4, 1),
        ).first()

        self.assertIsNotNone(april_due)
        self.assertEqual(april_due.amount, Decimal("1000.00"))


class ChrtCleanupSafetyTests(TestCase):
    def test_cleanup_preserves_non_chrt_historical_dues_for_canceled_plan(self):
        donor = User.objects.create_user(
            phone_number="+919000000007",
            name="Cleanup Safety Donor",
            password="secret",
        )
        pooja_option = PoojaOption.objects.create(code="R7", name="Recurring 7")
        day_option = PoojaDayOption.objects.create(
            code="REGSAFE",
            description="Regular",
            category=DayOptionCategory.CODE,
        )

        RecurringPoojaPlan.objects.create(
            donor=donor,
            pooja_option=pooja_option,
            day_option=day_option,
            recurrence_kind=RecurrenceKind.RECURRING,
            recurrence_frequency=RecurrenceFrequency.MONTHLY,
            start_date=date(2026, 1, 1),
            next_occurrence=None,
            amount=Decimal("1000.00"),
            is_active=False,
            pause_from=date(2026, 4, 1),
            pause_until=date.max,
            metadata={
                "canceled_at": "2026-04-03",
                "cancel_effective_from": "2026-04-01",
            },
        )

        for month_start in (date(2026, 1, 1), date(2026, 2, 1), date(2026, 3, 1)):
            PaymentRecord.objects.create(
                donor=donor,
                amount=Decimal("1000.00"),
                mode="pending",
                status=PaymentStatus.PENDING,
                payment_month=month_start,
                notes="Monthly recurring pooja contribution due",
                registration=None,
            )

        _clean_stale_chrt_dues(today=date(2026, 4, 3))

        remaining_months = list(
            PaymentRecord.objects.filter(
                donor=donor,
                status=PaymentStatus.PENDING,
                registration__isnull=True,
            )
            .order_by("payment_month")
            .values_list("payment_month", flat=True)
        )
        self.assertEqual(remaining_months, [date(2026, 1, 1), date(2026, 2, 1), date(2026, 3, 1)])

    def test_cleanup_does_not_normalize_historical_months_to_current_active_total(self):
        donor = User.objects.create_user(
            phone_number="+919000000008",
            name="Cleanup Normalize Safety Donor",
            password="secret",
        )
        day_option = PoojaDayOption.objects.create(
            code="REGSAFE2",
            description="Regular",
            category=DayOptionCategory.CODE,
        )

        # Active monthly plan (₹500)
        RecurringPoojaPlan.objects.create(
            donor=donor,
            pooja_option=PoojaOption.objects.create(code="R8A", name="Recurring 8 Active"),
            day_option=day_option,
            recurrence_kind=RecurrenceKind.RECURRING,
            recurrence_frequency=RecurrenceFrequency.MONTHLY,
            start_date=date(2026, 1, 1),
            next_occurrence=date(2026, 4, 1),
            amount=Decimal("500.00"),
            is_active=True,
        )

        # Canceled plan effective from Apr 2026 (₹2500)
        RecurringPoojaPlan.objects.create(
            donor=donor,
            pooja_option=PoojaOption.objects.create(code="R8C", name="Recurring 8 Canceled"),
            day_option=day_option,
            recurrence_kind=RecurrenceKind.RECURRING,
            recurrence_frequency=RecurrenceFrequency.MONTHLY,
            start_date=date(2026, 1, 1),
            next_occurrence=None,
            amount=Decimal("2500.00"),
            is_active=False,
            pause_from=date(2026, 4, 1),
            pause_until=date.max,
            metadata={
                "canceled_at": "2026-04-03",
                "cancel_effective_from": "2026-04-01",
            },
        )

        # Historical dues should stay ₹3000 and must not be normalized to active ₹500.
        for month_start in (date(2026, 1, 1), date(2026, 2, 1), date(2026, 3, 1)):
            PaymentRecord.objects.create(
                donor=donor,
                amount=Decimal("3000.00"),
                mode="pending",
                status=PaymentStatus.PENDING,
                payment_month=month_start,
                notes="Monthly recurring pooja contribution due",
                registration=None,
            )

        _clean_stale_chrt_dues(today=date(2026, 4, 3))

        amounts_by_month = list(
            PaymentRecord.objects.filter(
                donor=donor,
                status=PaymentStatus.PENDING,
                registration__isnull=True,
            )
            .order_by("payment_month")
            .values_list("payment_month", "amount")
        )
        self.assertEqual(
            amounts_by_month,
            [
                (date(2026, 1, 1), Decimal("3000.00")),
                (date(2026, 2, 1), Decimal("3000.00")),
                (date(2026, 3, 1), Decimal("3000.00")),
            ],
        )

    def test_cleanup_merges_duplicate_generic_month_rows_without_dropping_total(self):
        donor = User.objects.create_user(
            phone_number="+919000000009",
            name="Cleanup Duplicate Month Donor",
            password="secret",
        )
        day_option = PoojaDayOption.objects.create(
            code="REGSAFE3",
            description="Regular",
            category=DayOptionCategory.CODE,
        )
        RecurringPoojaPlan.objects.create(
            donor=donor,
            pooja_option=PoojaOption.objects.create(code="R9A", name="Recurring 9 Active"),
            day_option=day_option,
            recurrence_kind=RecurrenceKind.RECURRING,
            recurrence_frequency=RecurrenceFrequency.MONTHLY,
            start_date=date(2026, 1, 1),
            next_occurrence=date(2026, 4, 1),
            amount=Decimal("1000.00"),
            is_active=True,
        )

        # Legacy split rows for Jan 2026 (₹1000 + ₹1000) should be merged to ₹2000,
        # not collapsed to a single ₹1000 row.
        for amount in (Decimal("1000.00"), Decimal("1000.00")):
            PaymentRecord.objects.create(
                donor=donor,
                amount=amount,
                mode="pending",
                status=PaymentStatus.PENDING,
                payment_month=date(2026, 1, 1),
                notes="Monthly recurring pooja contribution due",
                registration=None,
            )

        _clean_stale_chrt_dues(today=date(2026, 4, 3))

        jan_rows = list(
            PaymentRecord.objects.filter(
                donor=donor,
                status=PaymentStatus.PENDING,
                registration__isnull=True,
                payment_month=date(2026, 1, 1),
            ).order_by("id")
        )
        self.assertEqual(len(jan_rows), 1)
        self.assertEqual(jan_rows[0].amount, Decimal("2000.00"))
        self.assertEqual(jan_rows[0].notes, "Monthly recurring pooja contribution due")

    def test_cleanup_merges_mixed_generic_and_chrt_noted_rows_without_amount_loss(self):
        donor = User.objects.create_user(
            phone_number="+919000000010",
            name="Cleanup Mixed Note Donor",
            password="secret",
        )
        day_option = PoojaDayOption.objects.create(
            code="REGSAFE4",
            description="Regular",
            category=DayOptionCategory.CODE,
        )
        RecurringPoojaPlan.objects.create(
            donor=donor,
            pooja_option=PoojaOption.objects.create(code="R10A", name="Recurring 10 Active"),
            day_option=day_option,
            recurrence_kind=RecurrenceKind.RECURRING,
            recurrence_frequency=RecurrenceFrequency.MONTHLY,
            start_date=date(2026, 1, 1),
            next_occurrence=date(2026, 4, 1),
            amount=Decimal("1000.00"),
            is_active=True,
        )

        PaymentRecord.objects.create(
            donor=donor,
            amount=Decimal("1000.00"),
            mode="pending",
            status=PaymentStatus.PENDING,
            payment_month=date(2026, 1, 1),
            notes="Monthly recurring pooja contribution due",
            registration=None,
        )
        PaymentRecord.objects.create(
            donor=donor,
            amount=Decimal("1000.00"),
            mode="pending",
            status=PaymentStatus.PENDING,
            payment_month=date(2026, 1, 1),
            notes="CHRT (Preferred Date) pooja contribution due",
            registration=None,
        )

        _clean_stale_chrt_dues(today=date(2026, 4, 3))

        jan_rows = list(
            PaymentRecord.objects.filter(
                donor=donor,
                status=PaymentStatus.PENDING,
                registration__isnull=True,
                payment_month=date(2026, 1, 1),
            ).order_by("id")
        )
        self.assertEqual(len(jan_rows), 1)
        self.assertEqual(jan_rows[0].amount, Decimal("2000.00"))
        self.assertEqual(jan_rows[0].notes, "Monthly recurring pooja contribution due")


class ExpenseRecordApiTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.url = reverse("expense-records-list")
        self.admin = User.objects.create_user(
            phone_number="+919100000001",
            name="Expense Admin",
            password="secret",
            role=UserRole.ADMIN,
            is_staff=True,
        )
        self.donor = User.objects.create_user(
            phone_number="+919100000002",
            name="Expense Donor",
            password="secret",
        )
        self.read_only_admin = User.objects.create_superuser(
            phone_number="+91 9999999998",
            name="Read Only Expense Admin",
            password="adminpass1",
        )
        self.hidden_expense_admin = User.objects.create_superuser(
            phone_number="+91 9999999997",
            name="Hidden Expense Admin",
            password="adminpass2",
        )

    def test_admin_can_create_expense_record(self):
        self.client.force_authenticate(self.admin)
        response = self.client.post(
            self.url,
            {
                "transaction_date": "2026-01-15",
                "category": "Maintenance",
                "amount": "1250.00",
                "transaction_no": "TXN-EXP-001",
                "comments": "Paid via bank transfer",
                "remarks": "Verified by admin",
            },
            format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        created = ExpenseRecord.objects.get(pk=response.json()["id"])
        self.assertEqual(created.created_by_id, self.admin.id)
        self.assertEqual(created.amount, Decimal("1250.00"))
        self.assertEqual(created.transaction_no, "TXN-EXP-001")
        self.assertEqual(created.comments, "Paid via bank transfer")
        self.assertEqual(created.remarks, "Verified by admin")

    def test_non_admin_cannot_list_or_create_expenses(self):
        self.client.force_authenticate(self.donor)
        list_response = self.client.get(self.url)
        self.assertEqual(list_response.status_code, status.HTTP_403_FORBIDDEN)

        create_response = self.client.post(
            self.url,
            {
                "transaction_date": "2026-01-15",
                "category": "Maintenance",
                "amount": "1250.00",
            },
            format="json",
        )
        self.assertEqual(create_response.status_code, status.HTTP_403_FORBIDDEN)

    def test_hidden_expense_admin_cannot_list_or_create_expenses(self):
        self.client.force_authenticate(self.hidden_expense_admin)
        list_response = self.client.get(self.url)
        self.assertEqual(list_response.status_code, status.HTTP_403_FORBIDDEN)
        self.assertEqual(
            list_response.json()["detail"],
            EXPENSE_TRACKER_ACCESS_DENIED_MESSAGE,
        )

        create_response = self.client.post(
            self.url,
            {
                "transaction_date": "2026-01-15",
                "category": "Maintenance",
                "amount": "1250.00",
            },
            format="json",
        )
        self.assertEqual(create_response.status_code, status.HTTP_403_FORBIDDEN)
        self.assertEqual(
            create_response.json()["detail"],
            EXPENSE_TRACKER_ACCESS_DENIED_MESSAGE,
        )

    def test_read_only_admin_can_list_expenses(self):
        ExpenseRecord.objects.create(
            transaction_date=date(2026, 1, 10),
            category="Maintenance",
            amount=Decimal("200.00"),
            created_by=self.admin,
        )
        self.client.force_authenticate(self.read_only_admin)
        response = self.client.get(self.url, {"month": "2026-01"})
        self.assertEqual(response.status_code, status.HTTP_200_OK)

    def test_monthly_list_returns_full_unpaginated_payload(self):
        ExpenseRecord.objects.create(
            transaction_date=date(2026, 1, 15),
            category="Category with metadata",
            amount=Decimal("10.00"),
            transaction_no="TXN-LIST-001",
            comments="List comments",
            remarks="List remarks",
            created_by=self.admin,
        )
        ExpenseRecord.objects.bulk_create(
            [
                ExpenseRecord(
                    transaction_date=date(2026, 1, 15),
                    category=f"Category {idx}",
                    amount=Decimal("10.00"),
                    created_by=self.admin,
                )
                for idx in range(24)
            ]
        )
        self.client.force_authenticate(self.admin)
        response = self.client.get(self.url, {"month": "2026-01"})
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        payload = response.json()
        self.assertIsInstance(payload, list)
        self.assertEqual(len(payload), 25)
        metadata_row = next((item for item in payload if item["transaction_no"] == "TXN-LIST-001"), None)
        self.assertIsNotNone(metadata_row)
        self.assertEqual(metadata_row["comments"], "List comments")
        self.assertEqual(metadata_row["remarks"], "List remarks")

    def test_admin_can_update_expense_record(self):
        self.client.force_authenticate(self.admin)
        expense = ExpenseRecord.objects.create(
            transaction_date=date(2026, 1, 15),
            category="Maintenance",
            amount=Decimal("1250.00"),
            transaction_no="TXN-EXP-OLD",
            comments="old comment",
            remarks="old remark",
            created_by=self.admin,
        )
        detail_url = reverse("expense-records-detail", args=[expense.id])
        response = self.client.patch(
            detail_url,
            {
                "category": "Updated Maintenance",
                "amount": "1500.00",
                "transaction_no": "TXN-EXP-NEW",
                "comments": "updated comment",
                "remarks": "updated remark",
            },
            format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        expense.refresh_from_db()
        self.assertEqual(expense.category, "Updated Maintenance")
        self.assertEqual(expense.amount, Decimal("1500.00"))
        self.assertEqual(expense.transaction_no, "TXN-EXP-NEW")
        self.assertEqual(expense.comments, "updated comment")
        self.assertEqual(expense.remarks, "updated remark")

    def test_admin_can_delete_expense_record(self):
        self.client.force_authenticate(self.admin)
        expense = ExpenseRecord.objects.create(
            transaction_date=date(2026, 1, 15),
            category="Maintenance",
            amount=Decimal("1250.00"),
            created_by=self.admin,
        )
        detail_url = reverse("expense-records-detail", args=[expense.id])
        response = self.client.delete(detail_url)
        self.assertEqual(response.status_code, status.HTTP_204_NO_CONTENT)
        self.assertFalse(ExpenseRecord.objects.filter(id=expense.id).exists())

    def test_read_only_admin_cannot_update_or_delete_expense(self):
        expense = ExpenseRecord.objects.create(
            transaction_date=date(2026, 1, 15),
            category="Maintenance",
            amount=Decimal("1250.00"),
            created_by=self.admin,
        )
        detail_url = reverse("expense-records-detail", args=[expense.id])
        self.client.force_authenticate(self.read_only_admin)

        update_response = self.client.patch(detail_url, {"category": "Blocked"}, format="json")
        self.assertEqual(update_response.status_code, status.HTTP_403_FORBIDDEN)

        delete_response = self.client.delete(detail_url)
        self.assertEqual(delete_response.status_code, status.HTTP_403_FORBIDDEN)
        self.assertTrue(ExpenseRecord.objects.filter(id=expense.id).exists())


class AdditionIncomeRecordApiTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.url = reverse("addition-income-records-list")
        self.admin = User.objects.create_user(
            phone_number="+919100000021",
            name="Addition Income Admin",
            password="secret",
            role=UserRole.ADMIN,
            is_staff=True,
        )
        self.donor = User.objects.create_user(
            phone_number="+919100000022",
            name="Addition Income Donor",
            password="secret",
        )
        self.read_only_admin = User.objects.create_superuser(
            phone_number="+91 9999999998",
            name="Read Only Addition Income Admin",
            password="adminpass1",
        )
        self.hidden_expense_admin = User.objects.create_superuser(
            phone_number="+91 9999999997",
            name="Hidden Addition Income Admin",
            password="adminpass2",
        )

    def test_admin_can_create_addition_income_record(self):
        self.client.force_authenticate(self.admin)
        response = self.client.post(
            self.url,
            {
                "transaction_date": "2026-01-18",
                "category": "General Donation",
                "amount": "2500.00",
                "transaction_no": "TXN-INC-001",
                "comments": "Received by cash",
                "remarks": "Verified by admin",
            },
            format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        created = AdditionIncomeRecord.objects.get(pk=response.json()["id"])
        self.assertEqual(created.created_by_id, self.admin.id)
        self.assertEqual(created.amount, Decimal("2500.00"))
        self.assertEqual(created.transaction_no, "TXN-INC-001")

    def test_non_admin_cannot_list_or_create_addition_income(self):
        self.client.force_authenticate(self.donor)
        list_response = self.client.get(self.url)
        self.assertEqual(list_response.status_code, status.HTTP_403_FORBIDDEN)

        create_response = self.client.post(
            self.url,
            {
                "transaction_date": "2026-01-18",
                "category": "General Donation",
                "amount": "2500.00",
            },
            format="json",
        )
        self.assertEqual(create_response.status_code, status.HTTP_403_FORBIDDEN)

    def test_hidden_admin_cannot_access_addition_income(self):
        self.client.force_authenticate(self.hidden_expense_admin)
        list_response = self.client.get(self.url)
        self.assertEqual(list_response.status_code, status.HTTP_403_FORBIDDEN)
        self.assertEqual(list_response.json()["detail"], EXPENSE_TRACKER_ACCESS_DENIED_MESSAGE)

    def test_admin_can_update_and_delete_addition_income_record(self):
        self.client.force_authenticate(self.admin)
        income = AdditionIncomeRecord.objects.create(
            transaction_date=date(2026, 1, 18),
            category="General Donation",
            amount=Decimal("2500.00"),
            transaction_no="TXN-INC-OLD",
            comments="old comment",
            remarks="old remark",
            created_by=self.admin,
        )
        detail_url = reverse("addition-income-records-detail", args=[income.id])
        update_response = self.client.patch(
            detail_url,
            {
                "category": "Updated Donation",
                "amount": "3000.00",
                "transaction_no": "TXN-INC-NEW",
            },
            format="json",
        )
        self.assertEqual(update_response.status_code, status.HTTP_200_OK)
        income.refresh_from_db()
        self.assertEqual(income.category, "Updated Donation")
        self.assertEqual(income.amount, Decimal("3000.00"))
        self.assertEqual(income.transaction_no, "TXN-INC-NEW")

        delete_response = self.client.delete(detail_url)
        self.assertEqual(delete_response.status_code, status.HTTP_204_NO_CONTENT)
        self.assertFalse(AdditionIncomeRecord.objects.filter(id=income.id).exists())

    def test_read_only_admin_can_list_but_cannot_modify_addition_income(self):
        income = AdditionIncomeRecord.objects.create(
            transaction_date=date(2026, 1, 18),
            category="General Donation",
            amount=Decimal("2500.00"),
            created_by=self.admin,
        )
        detail_url = reverse("addition-income-records-detail", args=[income.id])
        self.client.force_authenticate(self.read_only_admin)

        list_response = self.client.get(self.url, {"month": "2026-01"})
        self.assertEqual(list_response.status_code, status.HTTP_200_OK)

        update_response = self.client.patch(detail_url, {"category": "Blocked"}, format="json")
        self.assertEqual(update_response.status_code, status.HTTP_403_FORBIDDEN)

        delete_response = self.client.delete(detail_url)
        self.assertEqual(delete_response.status_code, status.HTTP_403_FORBIDDEN)
        self.assertTrue(AdditionIncomeRecord.objects.filter(id=income.id).exists())


class ExpenseCategoryApiTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.url = reverse("expense-categories-list")
        self.admin = User.objects.create_user(
            phone_number="+919100000101",
            name="Expense Category Admin",
            password="secret",
            role=UserRole.ADMIN,
            is_staff=True,
        )
        self.read_only_admin = User.objects.create_superuser(
            phone_number="+91 9999999998",
            name="Read Only Expense Category Admin",
            password="adminpass1",
        )
        self.hidden_expense_admin = User.objects.create_superuser(
            phone_number="+91 9999999997",
            name="Hidden Expense Category Admin",
            password="adminpass2",
        )

    def test_admin_can_create_and_filter_active_categories(self):
        self.client.force_authenticate(self.admin)
        create_response = self.client.post(
            self.url,
            {
                "name": "Electricity",
                "group_key": "coordinator",
                "display_order": 10,
                "is_active": True,
            },
            format="json",
        )
        self.assertEqual(create_response.status_code, status.HTTP_201_CREATED)
        ExpenseCategory.objects.create(
            name="Retired Item",
            group_key="other",
            display_order=99,
            is_active=False,
        )
        list_response = self.client.get(self.url, {"active": "true"})
        self.assertEqual(list_response.status_code, status.HTTP_200_OK)
        payload = list_response.json()
        names = {item["name"] for item in payload}
        self.assertIn("Electricity", names)
        self.assertNotIn("Retired Item", names)

    def test_read_only_admin_can_list_but_cannot_modify(self):
        category = ExpenseCategory.objects.create(name="Test Cat", group_key="other")
        detail_url = reverse("expense-categories-detail", args=[category.id])
        self.client.force_authenticate(self.read_only_admin)

        list_response = self.client.get(self.url)
        self.assertEqual(list_response.status_code, status.HTTP_200_OK)

        create_response = self.client.post(
            self.url,
            {"name": "Blocked Cat", "group_key": "other"},
            format="json",
        )
        self.assertEqual(create_response.status_code, status.HTTP_403_FORBIDDEN)

        update_response = self.client.patch(detail_url, {"name": "Blocked Update"}, format="json")
        self.assertEqual(update_response.status_code, status.HTTP_403_FORBIDDEN)

        delete_response = self.client.delete(detail_url)
        self.assertEqual(delete_response.status_code, status.HTTP_403_FORBIDDEN)

    def test_hidden_admin_cannot_access_expense_category_master(self):
        self.client.force_authenticate(self.hidden_expense_admin)
        list_response = self.client.get(self.url)
        self.assertEqual(list_response.status_code, status.HTTP_403_FORBIDDEN)
        self.assertEqual(list_response.json()["detail"], EXPENSE_TRACKER_ACCESS_DENIED_MESSAGE)


class IncomeCategoryApiTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.url = reverse("income-categories-list")
        self.admin = User.objects.create_user(
            phone_number="+919100000151",
            name="Income Category Admin",
            password="secret",
            role=UserRole.ADMIN,
            is_staff=True,
        )
        self.read_only_admin = User.objects.create_superuser(
            phone_number="+91 9999999918",
            name="Read Only Income Category Admin",
            password="adminpass1",
        )
        self.hidden_expense_admin = User.objects.create_superuser(
            phone_number="+91 9999999917",
            name="Hidden Income Category Admin",
            password="adminpass2",
        )

    def test_admin_can_create_and_list_income_categories(self):
        self.client.force_authenticate(self.admin)
        create_response = self.client.post(
            self.url,
            {
                "name": "General Donation",
                "display_order": 1,
            },
            format="json",
        )
        self.assertEqual(create_response.status_code, status.HTTP_201_CREATED)
        IncomeCategory.objects.create(name="Interest Income", display_order=2)

        list_response = self.client.get(self.url)
        self.assertEqual(list_response.status_code, status.HTTP_200_OK)
        payload = list_response.json()
        names = [item["name"] for item in payload]
        self.assertIn("General Donation", names)
        self.assertIn("Interest Income", names)

    def test_read_only_admin_can_list_but_cannot_modify(self):
        category = IncomeCategory.objects.create(name="Test Income Category")
        detail_url = reverse("income-categories-detail", args=[category.id])
        self.client.force_authenticate(self.read_only_admin)

        list_response = self.client.get(self.url)
        self.assertEqual(list_response.status_code, status.HTTP_200_OK)

        create_response = self.client.post(
            self.url,
            {"name": "Blocked Income Category"},
            format="json",
        )
        self.assertEqual(create_response.status_code, status.HTTP_403_FORBIDDEN)

        update_response = self.client.patch(detail_url, {"name": "Blocked Update"}, format="json")
        self.assertEqual(update_response.status_code, status.HTTP_403_FORBIDDEN)

        delete_response = self.client.delete(detail_url)
        self.assertEqual(delete_response.status_code, status.HTTP_403_FORBIDDEN)

    def test_hidden_admin_cannot_access_income_category_master(self):
        self.client.force_authenticate(self.hidden_expense_admin)
        list_response = self.client.get(self.url)
        self.assertEqual(list_response.status_code, status.HTTP_403_FORBIDDEN)
        self.assertEqual(list_response.json()["detail"], EXPENSE_TRACKER_ACCESS_DENIED_MESSAGE)


class PaymentRecordDeleteAccessTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.donor = User.objects.create_user(
            phone_number="+919200000001",
            name="Delete Donor",
            password="secret",
        )
        self.other_donor = User.objects.create_user(
            phone_number="+919200000002",
            name="Other Donor",
            password="secret",
        )
        self.admin = User.objects.create_user(
            phone_number="+919200000003",
            name="Delete Admin",
            password="secret",
            role=UserRole.ADMIN,
            is_staff=True,
        )

        self.success_payment = PaymentRecord.objects.create(
            donor=self.donor,
            amount=Decimal("250.00"),
            mode="upi",
            status=PaymentStatus.SUCCESS,
            transaction_reference="DEL-TXN-1",
            payment_month=date(2026, 2, 1),
        )
        self.pending_due = PaymentRecord.objects.create(
            donor=self.donor,
            amount=Decimal("400.00"),
            mode="cash",
            status=PaymentStatus.PENDING,
            payment_month=date(2026, 2, 1),
        )
        self.other_payment = PaymentRecord.objects.create(
            donor=self.other_donor,
            amount=Decimal("300.00"),
            mode="upi",
            status=PaymentStatus.SUCCESS,
            transaction_reference="DEL-TXN-2",
            payment_month=date(2026, 2, 1),
        )
        CombinePaymentMapping.objects.create(
            main_donor=self.donor,
            parent_donor=self.other_donor,
            effective_from=timezone.localdate().replace(day=1),
        )

    def _detail_url(self, payment_id: int) -> str:
        return reverse("payment-records-detail", args=[payment_id])

    def test_donor_without_access_cannot_delete_payment(self):
        profile = DonorProfile.objects.get(user=self.donor)
        profile.payment_delete_access = False
        profile.save(update_fields=["payment_delete_access"])
        self.client.force_authenticate(self.donor)

        response = self.client.delete(self._detail_url(self.success_payment.id))

        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)
        self.assertEqual(
            response.json()["detail"],
            "Please contact Admin for payment delete access.",
        )
        self.assertTrue(PaymentRecord.objects.filter(id=self.success_payment.id).exists())

    def test_donor_with_access_can_delete_own_success_payment(self):
        profile = DonorProfile.objects.get(user=self.donor)
        profile.payment_delete_access = True
        profile.save(update_fields=["payment_delete_access"])
        self.client.force_authenticate(self.donor)

        response = self.client.delete(self._detail_url(self.success_payment.id))

        self.assertEqual(response.status_code, status.HTTP_204_NO_CONTENT)
        self.assertFalse(PaymentRecord.objects.filter(id=self.success_payment.id).exists())

    def test_donor_with_access_cannot_delete_other_donor_payment(self):
        profile = DonorProfile.objects.get(user=self.donor)
        profile.payment_delete_access = True
        profile.save(update_fields=["payment_delete_access"])
        self.client.force_authenticate(self.donor)

        response = self.client.delete(self._detail_url(self.other_payment.id))

        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)
        self.assertEqual(
            response.json()["detail"],
            "You can delete only your own payment records.",
        )
        self.assertTrue(PaymentRecord.objects.filter(id=self.other_payment.id).exists())

    def test_donor_with_access_cannot_delete_pending_due_record(self):
        profile = DonorProfile.objects.get(user=self.donor)
        profile.payment_delete_access = True
        profile.save(update_fields=["payment_delete_access"])
        self.client.force_authenticate(self.donor)

        response = self.client.delete(self._detail_url(self.pending_due.id))

        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)
        self.assertEqual(
            response.json()["detail"],
            "Only received payment records can be deleted.",
        )
        self.assertTrue(PaymentRecord.objects.filter(id=self.pending_due.id).exists())

    def test_admin_can_delete_payment_without_donor_access(self):
        profile = DonorProfile.objects.get(user=self.donor)
        profile.payment_delete_access = False
        profile.save(update_fields=["payment_delete_access"])
        self.client.force_authenticate(self.admin)

        response = self.client.delete(self._detail_url(self.success_payment.id))

        self.assertEqual(response.status_code, status.HTTP_204_NO_CONTENT)
        self.assertFalse(PaymentRecord.objects.filter(id=self.success_payment.id).exists())
