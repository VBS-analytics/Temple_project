from decimal import Decimal, ROUND_HALF_UP, InvalidOperation
from pathlib import Path
from typing import Iterable

from django.core.management.base import BaseCommand, CommandError
from openpyxl import load_workbook

from ...models import DonorProfile, User


def _normalize_digits(value: str | None) -> str | None:
    if not value:
        return None
    digits = "".join(ch for ch in value if ch.isdigit())
    if not digits:
        return None
    return digits


def _header_label(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip().lower()


def _stringify(cell_value: object | None) -> str:
    if cell_value is None:
        return ""
    if isinstance(cell_value, float):
        if cell_value.is_integer():
            return str(int(cell_value))
        return str(cell_value)
    if isinstance(cell_value, Decimal):
        return format(cell_value, "f")
    return str(cell_value).strip()


def _parse_balance(value: object | None) -> int | None:
    text = _stringify(value)
    if text == "":
        return None
    try:
        decimal_value = Decimal(text)
    except InvalidOperation:
        return None
    rounded = decimal_value.to_integral_value(rounding=ROUND_HALF_UP)
    return int(rounded)


def _get_column_index(header_row: Iterable[object], column_name: str) -> int:
    normalized_target = column_name.strip().lower()
    for idx, cell_value in enumerate(header_row):
        if _header_label(cell_value) == normalized_target:
            return idx
    raise CommandError(f"Column '{column_name}' not found in the spreadsheet")


class Command(BaseCommand):
    help = "Import donor opening balances from an Excel spreadsheet."

    def add_arguments(self, parser):
        parser.add_argument("spreadsheet", type=str, help="Path to the XLSX file with opening balances.")
        parser.add_argument(
            "--sheet",
            type=str,
            help="Optional sheet name to read from (defaults to the active sheet).",
        )
        parser.add_argument("--name-column", type=str, default="Name", help="Column label containing donor names.")
        parser.add_argument("--phone-column", type=str, default="Phone", help="Column label containing phone numbers.")
        parser.add_argument(
            "--balance-column",
            type=str,
            default="Opening Balance",
            help="Column label containing the opening balance values.",
        )
        parser.add_argument("--dry-run", action="store_true", help="Parse the file and show the summary without saving.")

    def handle(self, *args, **options):
        spreadsheet_path = Path(options["spreadsheet"])
        if not spreadsheet_path.exists():
            raise CommandError(f"Spreadsheet not found at {spreadsheet_path}")

        workbook = load_workbook(spreadsheet_path, data_only=True)
        try:
            sheet_name = options.get("sheet")
            if sheet_name:
                if sheet_name not in workbook.sheetnames:
                    raise CommandError(f"Sheet '{sheet_name}' is not present in the workbook")
                sheet = workbook[sheet_name]
            else:
                sheet = workbook.active

            rows = sheet.iter_rows(values_only=True)
            header = next(rows, None)
            if not header:
                raise CommandError("The spreadsheet is empty")

            name_column_index = _get_column_index(header, options["name_column"])
            phone_column_index = _get_column_index(header, options["phone_column"])
            balance_column_index = _get_column_index(header, options["balance_column"])

            users = User.objects.all()
            digits_map: dict[str, list[User]] = {}
            suffix_map: dict[str, list[User]] = {}
            for user in users:
                normalized = _normalize_digits(user.phone_number)
                if not normalized:
                    continue
                digits_map.setdefault(normalized, []).append(user)
                if len(normalized) >= 10:
                    suffix_map.setdefault(normalized[-10:], []).append(user)

            processed_rows = 0
            updated_profiles = 0
            unchanged = 0
            unmatched_rows: list[dict[str, str | int]] = []

            for row_index, row in enumerate(rows, start=2):
                if not row or all(cell is None for cell in row):
                    continue
                processed_rows += 1
                name_value = row[name_column_index] if name_column_index < len(row) else None
                phone_value = row[phone_column_index] if phone_column_index < len(row) else None
                balance_value = row[balance_column_index] if balance_column_index < len(row) else None

                opening_balance = _parse_balance(balance_value)
                if opening_balance is None:
                    continue

                phone_digits = _normalize_digits(_stringify(phone_value))
                candidate_users: list[User] = []
                if phone_digits:
                    candidate_users = digits_map.get(phone_digits, [])
                    if not candidate_users and len(phone_digits) >= 10:
                        candidate_users = suffix_map.get(phone_digits[-10:], [])

                matched_user = None
                if len(candidate_users) == 1:
                    matched_user = candidate_users[0]
                elif len(candidate_users) > 1:
                    normalized_name = _stringify(name_value).replace(" ", "").lower()
                    for candidate in candidate_users:
                        if normalized_name and candidate.name.replace(" ", "").lower() == normalized_name:
                            matched_user = candidate
                            break
                if not matched_user:
                    unmatched_rows.append(
                        {
                            "row": row_index,
                            "name": _stringify(name_value),
                            "phone": _stringify(phone_value),
                        }
                    )
                    continue

                profile, _ = DonorProfile.objects.get_or_create(user=matched_user)
                previous_value = profile.custom_number
                if previous_value == opening_balance:
                    unchanged += 1
                    continue

                if not options["dry_run"]:
                    profile.custom_number = opening_balance
                    profile.save(update_fields=["custom_number"])
                updated_profiles += 1
        finally:
            workbook.close()

        self.stdout.write(self.style.SUCCESS(f"Parsed {processed_rows} rows from '{spreadsheet_path.name}'."))
        self.stdout.write(self.style.SUCCESS(f"Matched and updated {updated_profiles} donor profiles."))
        if unchanged:
            self.stdout.write(f"{unchanged} profiles already had the same balance.")
        if unmatched_rows:
            self.stdout.write(self.style.WARNING(f"Unable to match {len(unmatched_rows)} rows to donors:"))
            for unmatched in unmatched_rows:
                self.stdout.write(
                    self.style.WARNING(
                        f"Row {unmatched['row']} | Name: {unmatched['name'] or '-'} | Phone: {unmatched['phone'] or '-'}"
                    )
                )
        if options["dry_run"]:
            self.stdout.write(self.style.WARNING("Dry run enabled; no profiles were saved."))
